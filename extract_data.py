import os
import sys
import re
import json
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def get_gender(name, nick, all_text, intro_text):
    # Known overrides / specific mappings for ambiguous cases
    known_gender = {
        'สมัชญา คุ้มวงศ์': 'หญิง',
        'ณิชกานต์ ส้มเกลี้ยง': 'หญิง',
        'สุทธิภาสน์ เหลืองภัทรวงศ์': 'ชาย',
        'ธาราทิพย์ นาคปานเสือ': 'หญิง',
        'สิริณัฏฐ์ เนื่องจำนงค์': 'ชาย',
        'สิริชัย ศรีจงใจ': 'ชาย',
        'ณัฐณิชา แทนสอน': 'หญิง',
        'ชัยฟิตา ยูโซะ': 'หญิง',
        'พัชราภา ตาดไธสง': 'หญิง',
        'ชัญญานุช สายยาโน': 'หญิง',
        'วราภรณ์ ศรีปัญญา': 'หญิง',
        'ธนะเทพ ตะพัง': 'ชาย',
        'พรสินี วังคำแหง': 'หญิง',
        'พิชามญชุ์  เที่ยงจิตต์': 'หญิง',
        'พิชามญชุ์ เที่ยงจิตต์': 'หญิง',
        'ฐานนท์ โออินทร์': 'ชาย',
        'ณัฐสิริ บัวเจริญ': 'หญิง',
        'จินดาหรา บุตรธรรม': 'หญิง',
        'กานต์พิชชา จึงจีระสิทธิ์': 'หญิง',
        'พรรวินท์ หมื่นสุวรรณ์': 'หญิง',
        'ปัณณวิชญ์ ไชยเเก้ว': 'ชาย',
        'สิริกร กล่อมดี': 'หญิง',
        'กรกฎ สอนง่าย': 'หญิง',
        'รมิตา วรวงศ์': 'หญิง',
        'ชิษณุพงศ์ วงศ์เทพ': 'ชาย',
        'กัลยภรณ์ สุธารส': 'หญิง',
        'นันทัชพร รัตนะวัน': 'หญิง',
        'รสกร  สาระญาณ': 'หญิง',
        'รสกร สาระญาณ': 'หญิง',
        'ภากมล สุอุทัย': 'หญิง',
        'วรกมล ศรีคำม้วน': 'หญิง',
        'พศุตม์คเณศ ยอดวศิน': 'ชาย',
        'ศรเทพ รอดบัวทอง': 'ชาย',
        'ญาดา ไพโรจน์ฤทธิ์กุล': 'หญิง',
        'กอหญ้า บุญเทศนา': 'หญิง',
        'ธนดล เหล่าเคน': 'ชาย',
        'ธนภัทร ดีจริงตระกูล': 'ชาย',
        'ณัฏฐชัย แก้ววงศ์วัฒนา': 'ชาย',
        'ณัฐญดา เอื้อตระกูลวิทย์': 'หญิง',
        'ชยาภัสร์ ปุณโยทัย': 'หญิง',
        'เตชินท์ สุขสว่าง': 'ชาย',
        'ญาณิศา พรสุวรรณ': 'หญิง',
        'ภคมน วรจินดา': 'หญิง',
        'ทรงพล พรรณปราโมทย์': 'ชาย',
        'วันใหม่ เวชพันธ์': 'หญิง',
        'ปัณฑาลีย์ ธนัฐพัสวี': 'หญิง',
        'อรวรา เอี่ยมภิรมย์': 'หญิง',
        'ฐิติวรดา มุ่งอุ่น': 'หญิง',
        'ลธิชา ทวีกุล': 'หญิง',
        'ณัฐวิภา ทองประไพ': 'หญิง',
        'ชยานันท์ แสงพึ่งธรรม': 'หญิง',
        'สุกฤษฏิ์ นุชิต': 'ชาย',
        'วชิรพงศ์ ภูห้องเพชร': 'ชาย',
        'ธนวัฒน์ แสวงทอง': 'ชาย',
        'อุชุกร เจียมตน': 'ชาย',
        'ชยานันต์ เรืองเทพ': 'ชาย',
        'วัชรากร สากะจาย': 'ชาย',
        'อภิชาติ ศรีสุธรรม': 'ชาย',
        'ชินานันท์ ธนะสุชาวิรัตน์': 'หญิง',
        'ณัฐวัฒน์ พันธุ์พานิช': 'ชาย',
        'ณัฐมล อัมพรสินธุ์': 'หญิง',
        'พีรภัทร แสงสุรินทร์': 'ชาย',
    }

    clean_name = name.strip()
    if clean_name in known_gender:
        return known_gender[clean_name]

    # Check intro pronouns
    m_intro = len(re.findall(r'(ครับ|ผม\b|นาย\b|ฮะ|คราบ|คับ)', intro_text))
    f_intro = len(re.findall(r'(ค่ะ|คะ|หนู\b|ดิฉัน|นางสาว|ค่าา)', intro_text))
    if m_intro > f_intro:
        return 'ชาย'
    if f_intro > m_intro:
        return 'หญิง'

    # Check all text pronouns
    m_all = len(re.findall(r'(ครับ|ผม|ฮะ|คราบ|คับ)', all_text))
    f_all = len(re.findall(r'(ค่ะ|คะ|หนู|ดิฉัน|ค่าา|นะค้า)', all_text))
    if m_all > f_all:
        return 'ชาย'
    if f_all > m_all:
        return 'หญิง'

    # Feminine suffixes in Thai names
    fem_suffixes = ['พร', 'ภรณ์', 'วดี', 'รัตน์', 'กาญจน์', 'สุดา', 'นิภา', 'ณิชา', 'ศิริ', 'นภา', 'วรรณ', 'วัลย์', 'ดา', 'มาศ', 'เพ็ญ', 'ลดา']
    for s in fem_suffixes:
        if clean_name.endswith(s):
            return 'หญิง'

    return 'ไม่ระบุ'

def extract():
    excel_path = 'ค่ายวิศวะเหลาดินสอ ครั้งที่ 21 (การตอบกลับ).xlsx'
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['การตอบแบบฟอร์ม 1']

    row1 = next(ws.iter_rows(max_row=1, values_only=True))
    headers = [str(c).strip() if c is not None else '' for c in row1]

    # Department question ranges
    choice1_map = {
        'ฝ่ายพี่บ้าน': range(17, 23),
        'ฝ่ายกิจกรรม': range(23, 29),
        'ฝ่ายสถานที่': range(29, 35),
        'ฝ่ายสวัสดิการ': range(35, 41),
        'ฝ่ายพยาบาล': range(41, 47),
        'ฝ่ายทะเบียน': range(47, 53),
        'ฝ่ายสปอนเซอร์': range(53, 59),
        'ฝ่ายศิลป์': range(59, 65),
        'ฝ่ายประชาสัมพันธ์': range(65, 71),
        'ฝ่ายรันสคริป์': range(71, 76),
        'ฝ่ายโสตทัศนศึกษา': range(76, 82),
    }

    choice2_map = {
        'ฝ่ายพี่บ้าน': range(83, 89),
        'ฝ่ายกิจกรรม': range(89, 95),
        'ฝ่ายสถานที่': range(95, 101),
        'ฝ่ายสวัสดิการ': range(101, 107),
        'ฝ่ายพยาบาล': range(107, 113),
        'ฝ่ายทะเบียน': range(113, 119),
        'ฝ่ายสปอนเซอร์': range(119, 125),
        'ฝ่ายศิลป์': range(125, 131),
        'ฝ่ายประชาสัมพันธ์': range(131, 137),
        'ฝ่ายรันสคริป์': range(137, 142),
        'ฝ่ายโสตทัศนศึกษา': range(142, 148),
    }

    # General questions in columns 10 to 15 (1-based)
    general_q_indices = [10, 11, 12, 13, 14, 15]

    applicants = []
    
    # Specific people marked in red or yellow
    # From sheet inspection:
    red_names = {'ชินานันท์ ธนะสุชาวิรัตน์', 'ฐานนท์ โออินทร์', 'พีรภัทร แสงสุรินทร์', 'ณัฐวัฒน์ พันธุ์พานิช', 'ณัฐมล อัมพรสินธุ์'}
    yellow_names = {'ฐานนท์ โออินทร์', 'พีรภัทร แสงสุรินทร์'}

    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(row=r, column=3).value or '').strip()
        if not name:
            continue

        nick = str(ws.cell(row=r, column=4).value or '').strip()
        sid_raw = str(ws.cell(row=r, column=5).value or '').strip()
        if sid_raw.endswith('.0'):
            sid_raw = sid_raw[:-2]
        
        year = sid_raw[:2] if len(sid_raw) >= 2 else '??'
        major = str(ws.cell(row=r, column=6).value or '').strip()
        campus = str(ws.cell(row=r, column=7).value or '').strip()
        phone = str(ws.cell(row=r, column=8).value or '').strip()
        contact = str(ws.cell(row=r, column=9).value or '').strip()
        timestamp = str(ws.cell(row=r, column=1).value or '').strip()

        # Gather text for gender and searching
        all_row_text = ' '.join(str(ws.cell(row=r, column=c).value or '') for c in range(1, ws.max_column + 1))
        intro_text = str(ws.cell(row=r, column=10).value or '')

        gender = get_gender(name, nick, all_row_text, intro_text)

        # Department choices
        f1 = str(ws.cell(row=r, column=16).value or '').strip()
        f2 = str(ws.cell(row=r, column=82).value or '').strip()
        wants_f2_raw = str(ws.cell(row=r, column=150).value or '').strip()

        # General questions
        general_qa = []
        for q_idx in general_q_indices:
            q_text = headers[q_idx - 1]
            a_text = str(ws.cell(row=r, column=q_idx).value or '').strip()
            general_qa.append({
                'id': q_idx,
                'question': q_text,
                'answer': a_text
            })

        # Choice 1 questions
        choice1_qa = []
        if f1 in choice1_map:
            for q_idx in choice1_map[f1]:
                if q_idx <= len(headers):
                    q_text = headers[q_idx - 1]
                    a_text = str(ws.cell(row=r, column=q_idx).value or '').strip()
                    choice1_qa.append({
                        'id': q_idx,
                        'question': q_text,
                        'answer': a_text
                    })

        # Choice 2 questions
        choice2_qa = []
        if f2 in choice2_map:
            for q_idx in choice2_map[f2]:
                if q_idx <= len(headers):
                    q_text = headers[q_idx - 1]
                    # strip ' 2' suffix from title if present for clean display
                    clean_q_text = re.sub(r'\s*2\s*$', '', q_text).strip()
                    a_text = str(ws.cell(row=r, column=q_idx).value or '').strip()
                    choice2_qa.append({
                        'id': q_idx,
                        'question': clean_q_text,
                        'raw_question': q_text,
                        'answer': a_text
                    })

        # Interview preference
        interview_pref = str(ws.cell(row=r, column=148).value or '').strip()

        # Exclusion checks
        is_red = name in red_names
        is_yellow = name in yellow_names
        is_year67 = (year == '67')

        exclusion_reasons = []
        if is_red:
            exclusion_reasons.append('ชื่อแดง')
        if is_yellow:
            exclusion_reasons.append('ชื่อเหลือง')
        if is_year67:
            exclusion_reasons.append('รหัสนักศึกษา 67')

        is_excluded = len(exclusion_reasons) > 0

        # Experience check
        exp_text = str(ws.cell(row=r, column=11).value or '').strip().lower()
        has_experience = any(k in exp_text for k in ['เคยทำ', 'เคยทำค่าย', 'เคยเป็น', 'เคย']) and not any(k in exp_text for k in ['ไม่เคย', 'ยังไม่เคย', 'ไม่เคยทำ'])

        applicant_data = {
            'id': r - 1,
            'row': r,
            'timestamp': timestamp,
            'name': name,
            'nickname': nick,
            'studentId': sid_raw,
            'year': year,
            'major': major,
            'campus': campus,
            'phone': phone,
            'contact': contact,
            'gender': gender,
            'choice1': {
                'dept': f1,
                'questions': choice1_qa
            },
            'choice2': {
                'dept': f2 if f2 else 'ไม่ได้เลือกอันดับ 2',
                'questions': choice2_qa
            },
            'generalQuestions': general_qa,
            'interviewPref': interview_pref,
            'wantsChoice2': bool(f2 and f2 != 'ไม่เลือก' and f2 != 'ไม่ได้เลือก'),
            'hasExperience': has_experience,
            'isExcluded': is_excluded,
            'exclusionReasons': exclusion_reasons,
            'isRed': is_red,
            'isYellow': is_yellow,
            'isYear67': is_year67,
            'searchIndex': f"{name} {nick} {sid_raw} {major} {f1} {f2} {gender}".lower()
        }
        applicants.append(applicant_data)

    # Compute overall statistics (both all and eligible)
    eligible_applicants = [a for a in applicants if not a['isExcluded']]

    def get_stats_for_list(alist):
        dept_c1 = {}
        dept_c2 = {}
        dept_total = {}
        year_dist = {}
        gender_dist = {}
        major_dist = {}
        interview_dist = {}
        exp_dist = {'เคยทำค่าย': 0, 'ไม่เคยทำค่าย': 0}

        all_depts = [
            'ฝ่ายพี่บ้าน', 'ฝ่ายกิจกรรม', 'ฝ่ายสถานที่', 'ฝ่ายสวัสดิการ',
            'ฝ่ายพยาบาล', 'ฝ่ายทะเบียน', 'ฝ่ายสปอนเซอร์', 'ฝ่ายศิลป์',
            'ฝ่ายประชาสัมพันธ์', 'ฝ่ายรันสคริป์', 'ฝ่ายโสตทัศนศึกษา'
        ]
        for d in all_depts:
            dept_c1[d] = 0
            dept_c2[d] = 0
            dept_total[d] = 0

        for a in alist:
            f1 = a['choice1']['dept']
            f2 = a['choice2']['dept']
            if f1 in dept_c1:
                dept_c1[f1] += 1
                dept_total[f1] += 1
            if f2 in dept_c2:
                dept_c2[f2] += 1
                dept_total[f2] += 1

            y = a['year']
            year_dist[y] = year_dist.get(y, 0) + 1

            g = a['gender']
            gender_dist[g] = gender_dist.get(g, 0) + 1

            m = a['major']
            major_dist[m] = major_dist.get(m, 0) + 1

            pref = a['interviewPref']
            if '10' in pref and '11' in pref:
                p_label = 'สะดวกทั้ง 10 และ 11 ต.ค.'
            elif '10' in pref:
                p_label = 'สะดวก 10 ต.ค.'
            elif '11' in pref:
                p_label = 'สะดวก 11 ต.ค.'
            else:
                p_label = 'อื่นๆ / ไม่ระบุ'
            interview_dist[p_label] = interview_dist.get(p_label, 0) + 1

            if a['hasExperience']:
                exp_dist['เคยทำค่าย'] += 1
            else:
                exp_dist['ไม่เคยทำค่าย'] += 1

        return {
            'total': len(alist),
            'deptChoice1': dept_c1,
            'deptChoice2': dept_c2,
            'deptTotal': dept_total,
            'yearDist': year_dist,
            'genderDist': gender_dist,
            'majorDist': major_dist,
            'interviewDist': interview_dist,
            'expDist': exp_dist
        }

    dataset = {
        'campName': 'ค่ายวิศวะเหลาดินสอ ครั้งที่ 21',
        'generatedAt': '2026-10-06T17:40:00+07:00',
        'totalCount': len(applicants),
        'eligibleCount': len(eligible_applicants),
        'excludedCount': len(applicants) - len(eligible_applicants),
        'excludedList': [
            {
                'name': a['name'],
                'nickname': a['nickname'],
                'studentId': a['studentId'],
                'year': a['year'],
                'major': a['major'],
                'choice1': a['choice1']['dept'],
                'choice2': a['choice2']['dept'],
                'reasons': a['exclusionReasons']
            }
            for a in applicants if a['isExcluded']
        ],
        'departments': [
            'ฝ่ายพี่บ้าน', 'ฝ่ายกิจกรรม', 'ฝ่ายสถานที่', 'ฝ่ายสวัสดิการ',
            'ฝ่ายพยาบาล', 'ฝ่ายทะเบียน', 'ฝ่ายสปอนเซอร์', 'ฝ่ายศิลป์',
            'ฝ่ายประชาสัมพันธ์', 'ฝ่ายรันสคริป์', 'ฝ่ายโสตทัศนศึกษา'
        ],
        'stats': {
            'all': get_stats_for_list(applicants),
            'eligible': get_stats_for_list(eligible_applicants)
        },
        'applicants': applicants
    }

    with open('data.json', 'w', encoding='utf-8') as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)

    # Also export data.js for seamless offline / file:// loading
    with open('data.js', 'w', encoding='utf-8') as f:
        f.write('window.INITIAL_DATA = ' + json.dumps(dataset, ensure_ascii=False) + ';\n')

    print(f"Successfully processed {len(applicants)} applicants.")
    print(f"Eligible applicants: {len(eligible_applicants)}")
    print(f"Excluded applicants: {len(applicants) - len(eligible_applicants)}")
    for ex in dataset['excludedList']:
        print(f"  - {ex['name']} ({ex['nickname']}) ID:{ex['studentId']} -> {', '.join(ex['reasons'])}")

if __name__ == '__main__':
    extract()
